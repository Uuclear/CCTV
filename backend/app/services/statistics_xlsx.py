"""Export project segment table — aligned with 公共通道统计表.xlsx CCTV sheet."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.grading import normalize_condition_grade
from app.models import Defect, Segment

CCTV_HEADERS = [
    "序号",
    "管道",
    "管材",
    "管段",
    "管径（mm）",
    "长度（m）",
    "缺陷",
    "修复指数 RI",
    "结构状况",
    "养护指数 MI",
    "功能状况",
    "备注",
]


def _template_path(repo_root: Path) -> Path | None:
    for rel in (
        "公共通道统计表.xlsx",
        "templates/公共通道统计表.xlsx",
        "templates/export/公共通道统计表.xlsx",
    ):
        p = repo_root / rel
        if p.is_file():
            return p
    return None


def _defect_summary(defects: list[Defect]) -> str:
    if not defects:
        return "无"
    return "；".join(f"{d.defect_code}·L{d.level}" for d in defects[:8])


async def build_statistics_xlsx(
    db: AsyncSession,
    project_id: int,
    repo_root: Path,
) -> bytes:
    r = await db.execute(
        select(Segment).where(Segment.project_id == project_id).order_by(Segment.id)
    )
    segs = list(r.scalars().all())
    seg_ids = [s.id for s in segs]
    by_seg: dict[int, list[Defect]] = {}
    if seg_ids:
        dr = await db.execute(select(Defect).where(Defect.segment_id.in_(seg_ids)))
        for d in dr.scalars().all():
            by_seg.setdefault(d.segment_id, []).append(d)

    tpl = _template_path(repo_root)
    if tpl:
        wb = load_workbook(tpl)
        ws = wb["CCTV"] if "CCTV" in wb.sheetnames else wb.active
        start_row = 2
        # clear old data rows (keep header row 1)
        if ws.max_row >= start_row:
            ws.delete_rows(start_row, ws.max_row - start_row + 1)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "CCTV"
        for col, h in enumerate(CCTV_HEADERS, start=1):
            ws.cell(row=1, column=col, value=h)
        start_row = 2

    for i, seg in enumerate(segs, start=1):
        row = start_row + i - 1
        defects = by_seg.get(seg.id, [])
        a = seg.chain_start_label or "?"
        b = seg.chain_end_label or "?"
        ri = seg.ri if seg.ri is not None else ""
        mi = seg.mi if seg.mi is not None else ""
        values = [
            i,
            seg.pipe_system or "",
            seg.pipe_material or "",
            f"{a}～{b}",
            seg.diameter_mm if seg.diameter_mm is not None else "",
            seg.pipe_length_m if seg.pipe_length_m is not None else "",
            _defect_summary(defects),
            ri,
            normalize_condition_grade(seg.ri_grade, seg.ri),
            mi,
            normalize_condition_grade(seg.mi_grade, seg.mi),
            seg.remark or "",
        ]
        for col, val in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=val)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
